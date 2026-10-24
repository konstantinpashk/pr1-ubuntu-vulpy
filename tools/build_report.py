#!/usr/bin/env python3
"""Создание DOCX отчёта по подтверждённому ручному анализу и, если есть, Bandit JSON summary."""
from pathlib import Path
import argparse, csv, json
from datetime import date
from collections import Counter
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

P=argparse.ArgumentParser();P.add_argument('--input', required=True);P.add_argument('--summary');P.add_argument('--out',required=True);a=P.parse_args()
with Path(a.input).open(encoding='utf-8-sig',newline='') as f: findings=list(csv.DictReader(f))
summary=None
if a.summary and Path(a.summary).exists():
    summary=json.loads(Path(a.summary).read_text(encoding='utf-8'))

DOC=Document();sec=DOC.sections[0]
sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.2);sec.right_margin=Cm(1.9)
styles=DOC.styles
normal=styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(9.3);normal.font.color.rgb=RGBColor(37,49,67)
normal.paragraph_format.space_after=Pt(5)
for name,size in [('Title',24),('Heading 1',14),('Heading 2',11)]:
    st=styles[name];st.font.name='DejaVu Sans';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(25,54,89)
    st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(7)

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fr=footer.add_run('ПР1 · Ubuntu · Vulpy · статический анализ');fr.font.name='DejaVu Sans';fr.font.size=Pt(8);fr.font.color.rgb=RGBColor(108,120,136)

p=DOC.add_paragraph(style='Title');p.add_run('ПРАКТИЧЕСКАЯ РАБОТА № 1')
p=DOC.add_paragraph();r=p.add_run('GitHub и оценка безопасности исходного кода в Ubuntu');r.bold=True;r.font.size=Pt(14)
DOC.add_paragraph('Вариант № 1 — анализ исходного кода проекта',style='Subtitle')
DOC.add_paragraph('Обучающийся: _______________________________________')
DOC.add_paragraph('Группа: ____________________________________________')
DOC.add_paragraph('Преподаватель: _____________________________________')
DOC.add_paragraph('Дата подготовки: 24 октября 2026 г. (обновить на дату сдачи при необходимости)')

DOC.add_heading('1. Цель и задачи',1)
DOC.add_paragraph('Цель: освоить Git и GitHub из командной строки Bash на Ubuntu, выполнить анализ безопасности учебного проекта средствами статического анализа и представить машиночитаемые датасеты с обоснованной экспертной оценкой.')
DOC.add_paragraph('Задачи: (1) подготовить Ubuntu, Git и Python; (2) получить проект из открытого GitHub-репозитория; (3) проверить исходный код, сохранив формальные результаты; (4) разместить учебную копию проекта и отчёты в собственном GitHub-репозитории; (5) оценить риски и способы их уменьшения.')

DOC.add_heading('2. Объект и условия исследования',1)
DOC.add_paragraph('Проект: Vulpy (fportantier/vulpy), Python / Flask / SQLite. Публичный исходный репозиторий: https://github.com/fportantier/vulpy. Лицензия MIT. В проекте намеренно присутствуют две реализации: bad (уязвимая) и good (пример более безопасных решений; её безопасность не гарантирована).')
DOC.add_paragraph('Операционная система: Ubuntu; интерфейс: Bash; система контроля версий: Git; платформа публикации: GitHub. Инструмент SAST: Bandit. Проверяется исходный код без поднятия веб-сервера и без эксплуатации уязвимостей.')
DOC.add_paragraph('Снимок репозитория фиксируется git commit SHA в файле metadata/upstream_commit.txt после выполнения скрипта подготовки. Публикуемый репозиторий студента и реальная дата запусков заполняются по фактическим действиям.')
upstream=Path('metadata/upstream_commit.txt')
if upstream.exists():
    DOC.add_paragraph('Фактический upstream SHA: '+upstream.read_text(encoding='utf-8').strip())
else:
    DOC.add_paragraph('Фактический upstream SHA: ожидается выполнение 01_setup_ubuntu.sh на Ubuntu.')

DOC.add_heading('3. Методика и воспроизводимость',1)
for i,t in enumerate([
    'Получить код командой git clone на Ubuntu, отделить исходный каталог от вложенного .git и сохранить SHA источника.',
    'Создать Python venv и установить Bandit отдельно от уязвимых зависимостей приложения.',
    'Проанализировать bad/ и good/ с помощью Bandit, сформировать JSON-отчёты, CSV с находками и summary.json.',
    'Сопоставить автоматические предупреждения с ручной экспертизой исходников; оценить уровень риска и способы исправления.',
    'Выполнить git add/commit/push в собственный приватный GitHub-репозиторий, подтвердить историю коммитов и URL.'
],1):DOC.add_paragraph(f'{i}. {t}')

DOC.add_heading('4. Результаты ручного анализа исходников',1)
DOC.add_paragraph(f'Сформировано {len(findings)} записей ручного анализа. Они основаны на просмотре публичного исходного кода и НЕ являются числом срабатываний Bandit. Экспертные оценки относительные, не представляют собой рассчитанный CVSS.')
counts=Counter(x['Оценка риска (экспертная)'] for x in findings)
DOC.add_paragraph('Распределение ручных записей: '+', '.join(f'{name} — {counts[name]}' for name in ('Высокий','Средний','Низкий') if counts.get(name))+'.')

hdrs=['ID','Выявленная проблема','Риск','Файл / строки']
tbl=DOC.add_table(rows=1, cols=4);tbl.style='Light Shading Accent 1';tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(hdrs):tbl.rows[0].cells[i].text=h
for f in findings:
    c=tbl.add_row().cells
    for i,v in enumerate([f['ID'],f['Находка'],f['Оценка риска (экспертная)'],f["Файл"].removeprefix('bad/')+':'+f['Строки']]):c[i].text=v
    for x in c:x.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
for row in tbl.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            for r in p.runs:r.font.size=Pt(7.8)

DOC.add_heading('5. Результаты автоматического сканирования Bandit',1)
if summary is None:
    p=DOC.add_paragraph('Статус: НЕ ВЫПОЛНЕНО в среде подготовки отчёта.');p.runs[0].bold=True
    DOC.add_paragraph('Для получения реальных цифр выполните на Ubuntu последовательно bash 01_setup_ubuntu.sh и bash 02_scan.sh. Второй скрипт создаст reports/bandit_bad.json, reports/bandit_good.json, datasets/bandit_bad.csv, datasets/bandit_good.csv и финальную версию данного DOCX.')
else:
    stamp=summary.get('generated_at_utc','нет даты')
    DOC.add_paragraph('Дата и время запуска Bandit (UTC): '+stamp)
    t=DOC.add_table(rows=1,cols=5);t.style='Light Shading Accent 1'
    for i,h in enumerate(['Каталог','Всего','HIGH','MEDIUM','LOW']):t.rows[0].cells[i].text=h
    for name in ['bad','good']:
        v=summary['profiles'][name];s=v['by_severity'];row=t.add_row().cells
        for i,x in enumerate([name,str(v['total']),str(s['HIGH']),str(s['MEDIUM']),str(s['LOW'])]):row[i].text=x
    DOC.add_paragraph('Все отдельные срабатывания и их line_number, test_id, severity, confidence доступны в оригинальном Bandit JSON и в датасете CSV. Числа относятся только к автоматическим находкам, а не к ручной экспертной таблице.')
    for name in ['bad','good']:
        errs=summary['profiles'][name].get('errors',[])
        if errs:DOC.add_paragraph(f'Замечания/ошибки анализатора {name}: {len(errs)} (проверить исходный JSON).')
DOC.add_paragraph('Ограничение: Bandit ищет известные небезопасные конструкции Python. Полноценный анализ XSS, CSRF, авторизации, уязвимостей зависимостей и логики приложения потребует дополнительных инструментов и проверки. Нулевое число находок не доказывает безопасность.')

DOC.add_heading('6. Экспертное заключение',1)
DOC.add_paragraph('В специально уязвимой версии bad присутствуют критически значимые с точки зрения защиты данных и пользователей конструкции: конкатенация в SQL-запросах, небезопасное хранение и обработка паролей, неподписанная cookie с идентификатором пользователя, а также вывод недоверенных данных без HTML-экранирования. При реальном развёртывании это создавало бы высокий риск нарушения конфиденциальности и целостности данных.')
DOC.add_paragraph('Итоговая качественная оценка: уровень риска bad — высокий, использование в продуктивной или доступной извне среде недопустимо. Для учебной практики допускается только изолированный статический анализ. Результаты автоматического сканера служат дополнительными свидетельствами, но не заменяют ручную проверку.')
DOC.add_paragraph('Рекомендации по приоритету: (1) параметризация всех SQL-запросов; (2) защищённое хеширование паролей и исключение трассировки секретов; (3) подписанные сеансы и строгий контроль доступа; (4) отказ от safe для пользовательского HTML; (5) CSRF-контроль и запрет изменяющих GET; (6) CSPRNG для секретов; (7) повторное SAST-сканирование после исправлений.')

DOC.add_heading('7. Набор датасетов и доказательств',1)
for line in [
  'datasets/Ручной_аудит_Vulpy.csv и .json — описания проблем с CWE, уровнем риска, строками и URL исходника.',
  'datasets/Ручной_аудит_Vulpy.xlsx — табличное представление экспертного датасета.',
  'reports/bandit_bad.json, bandit_good.json — первичные результаты Bandit (после запуска).',
  'datasets/bandit_bad.csv, bandit_good.csv — машиночитаемая таблица предупреждений Bandit (после запуска).',
  'reports/summary.json — сведения о количестве находок и дате выполнения (после запуска).',
  'metadata/upstream_commit.txt — SHA изученной версии (после загрузки кода).',
  'Скриншоты Ubuntu-терминала, GitHub URL и git log — преподаватель прикладывает студент после выполнения.'
]:DOC.add_paragraph('• '+line)
DOC.add_heading('8. Источники и оговорки',1)
DOC.add_paragraph('Репозиторий: https://github.com/fportantier/vulpy')
DOC.add_paragraph('Исходники для проверки: https://raw.githubusercontent.com/fportantier/vulpy/master/bad/')
DOC.add_paragraph('Bandit: https://bandit.readthedocs.io/ — статический анализ безопасности Python.')
DOC.add_paragraph('В работе не выполнялись действия в чужих системах: исходники открытые, целевой проект создан именно как лаборатория уязвимостей. Все указанные значения по Bandit появляются только из настоящего запуска сканера на Ubuntu.')

out=Path(a.out);out.parent.mkdir(parents=True,exist_ok=True);DOC.save(out)
print(f'DOCX: {out}; manual findings: {len(findings)}; Bandit summary: {bool(summary)}')
